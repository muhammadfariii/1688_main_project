import codecs
"""
Automated Test Suite for 1688 Sourcing Tool.
Verifies classifier, Apify adapter, search engine, exporter, config auto-detection, deduplication, and FastAPI endpoints.
"""

import os
import io
import unittest
from unittest.mock import patch, MagicMock
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from config import config, AppConfig
from classifier import classifier, HeuristicSupplierClassifier
from adapters import DemoAdapter, Apify1688Adapter, get_adapter
from search_engine import SourcingSearchEngine
from exporter import SourcingDataExporter
from app import app


class TestConfig(unittest.TestCase):
    def test_explicit_live_mode(self):
        with patch.dict(os.environ, {"PROVIDER_API_KEY": "apify_api_123456789", "DEMO_MODE": "false"}, clear=False):
            cfg = AppConfig()
            status = cfg.get_public_status()
            self.assertFalse(cfg.demo_mode)
            self.assertEqual(cfg.provider_name, "apify")
            self.assertTrue(status["has_api_key"])
            self.assertEqual(status["masked_api_key"], "apify_ap...6789")
            self.assertIn("Live Mode", status["status_label"])

    def test_default_demo_mode_safety(self):
        from pathlib import Path
        with patch("config.LOCAL_CONFIG_FILE", Path("/tmp/nonexistent.json")):
            with patch.dict(os.environ, {"PROVIDER_API_KEY": "apify_api_123"}, clear=False):
                cfg = AppConfig()
                self.assertTrue(cfg.demo_mode)
                self.assertEqual(cfg.get_public_status()["provider_name"], "demo")

    def test_runtime_toggle_preservation(self):
        from pathlib import Path
        with patch("config.LOCAL_CONFIG_FILE", Path("/tmp/mock_config.json")):
            cfg = AppConfig()
            cfg.set_demo_mode(True)
            self.assertTrue(cfg.demo_mode)
            cfg.reload()
            self.assertTrue(cfg.demo_mode)

    def test_toggle_demo_mode_without_api_key(self):
        cfg = AppConfig()
        cfg.provider_api_key = ""
        success = cfg.set_demo_mode(False)
        self.assertFalse(success)
        self.assertTrue(cfg.demo_mode)


class TestAdapters(unittest.TestCase):
    def test_get_adapter_demo_when_in_demo_mode(self):
        cfg = AppConfig()
        cfg.provider_api_key = ""
        cfg.demo_mode = True
        adapter = get_adapter(cfg)
        self.assertIsInstance(adapter, DemoAdapter)
        self.assertTrue(adapter.get_provider_info()["is_demo"])

    def test_get_adapter_apify_when_key_present(self):
        cfg = AppConfig()
        cfg.provider_api_key = "apify_api_test_key"
        cfg.demo_mode = False
        adapter = get_adapter(cfg)
        self.assertIsInstance(adapter, Apify1688Adapter)
        self.assertFalse(adapter.get_provider_info()["is_demo"])

    def test_apify_endpoint_resolution(self):
        adapter = Apify1688Adapter(api_token="apify_api_123", actor_id="schnellscrapers/1688-supplier-leads", base_url="https://api.apify.com/v2")
        self.assertEqual(adapter.clean_actor_id, "schnellscrapers~1688-supplier-leads")
        self.assertEqual(
            adapter._get_run_sync_url(),
            "https://api.apify.com/v2/acts/schnellscrapers~1688-supplier-leads/run-sync-get-dataset-items"
        )

    def test_apify_response_extraction_varieties(self):
        adapter = Apify1688Adapter(api_token="test_token")
        # Direct list
        self.assertEqual(len(adapter._extract_items_from_response([{"supplierName": "A"}])), 1)
        # items list
        self.assertEqual(len(adapter._extract_items_from_response({"items": [{"supplierName": "A"}]})), 1)
        # data list
        self.assertEqual(len(adapter._extract_items_from_response({"data": [{"supplierName": "A"}, {"supplierName": "B"}]})), 2)
        # Single supplier dict
        self.assertEqual(len(adapter._extract_items_from_response({"supplierId": "b2b-123", "supplierName": "Factory A"})), 1)

    def test_apify_test_connection_no_key(self):
        adapter = Apify1688Adapter(api_token="")
        res = adapter.test_connection()
        self.assertFalse(res["success"])
        self.assertIn("missing", res["error"].lower())

    @patch("requests.post")
    def test_apify_search_mocked_success(self, mock_post):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.json.return_value = [
            {
                "searchMatches": [{"keyword": "手机壳", "position": 1}],
                "supplierId": "b2b-2850655109d72ea",
                "supplierName": "佛山市南海区三丰手机配件有限公司",
                "supplierUrl": "http://shop1460393846166.1688.com/",
                "location": "广东, 佛山市",
                "businessRole": "生产加工",
                "factoryStatus": "super_factory",
                "yearsOnPlatform": 11,
                "isFactoryInspected": True,
                "isBusinessInspected": False,
                "isSuperFactory": True,
                "transactionCount": 2536,
                "representativeOffers": [
                    {
                        "offerId": "931027992821",
                        "title": "透明磁吸手机壳",
                        "url": "https://detail.1688.com/offer/931027992821.html",
                        "priceCny": 4.2,
                        "minimumOrderQuantity": 30,
                        "transactionCount": 2536
                    }
                ],
                "sourceKeywords": ["手机壳"]
            }
        ]
        mock_post.return_value = mock_resp

        adapter = Apify1688Adapter(api_token="apify_api_test")
        results = adapter.search_single_product("手机壳")
        self.assertEqual(len(results), 1)
        res = results[0]
        self.assertEqual(res["supplier_id"], "b2b-2850655109d72ea")
        self.assertEqual(res["supplier_name"], "佛山市南海区三丰手机配件有限公司")
        self.assertEqual(res["city"], "佛山市")
        self.assertEqual(res["primary_product"], "透明磁吸手机壳")
        self.assertEqual(res["price_range_rmb"], "¥4.20")
        self.assertEqual(res["moq"], "30 pcs")
        self.assertTrue(res["is_super_factory"])
        self.assertTrue(res["is_factory_inspected"])
        self.assertIn("超级工厂 (Super Factory)", res["badges"])
        self.assertIn("深度验厂 (Factory Inspected)", res["badges"])

    def test_url_builders(self):
        from adapters import build_1688_product_url, build_1688_shop_url, build_1688_contact_page_url, build_1688_chat_url
        self.assertEqual(
            build_1688_product_url(931027992821),
            "https://detail.1688.com/offer/931027992821.html"
        )
        self.assertEqual(
            build_1688_product_url(None, "http://detail.1688.com/offer/931027992821.html"),
            "https://detail.1688.com/offer/931027992821.html"
        )
        self.assertEqual(
            build_1688_shop_url("http://shop1460393846166.1688.com/"),
            "https://shop1460393846166.1688.com/"
        )
        self.assertEqual(
            build_1688_contact_page_url("https://shop1460393846166.1688.com/"),
            "https://shop1460393846166.1688.com/page/contactinfo.htm"
        )
        self.assertIn(
            "touid=b2b-2850655109d72ea",
            build_1688_chat_url("b2b-2850655109d72ea")
        )

    def test_apify_deduplication(self):
        adapter = Apify1688Adapter(api_token="apify_api_test")
        dup_items = [
            {
                "supplier_id": "b2b-1001",
                "supplier_name": "Yongkang Factory A",
                "shop_url": "https://shop1001.1688.com",
                "products_matched": ["Water Bottle Standard"],
                "primary_product": "Water Bottle Standard",
                "badges": ["超级工厂"]
            },
            {
                "supplier_id": "b2b-1001",
                "supplier_name": "Yongkang Factory A",
                "shop_url": "https://shop1001.1688.com",
                "products_matched": ["Vacuum Flask Custom"],
                "primary_product": "Vacuum Flask Custom",
                "badges": ["深度验厂"]
            }
        ]
        deduped = adapter._deduplicate_suppliers(dup_items)
        self.assertEqual(len(deduped), 1)
        self.assertEqual(len(deduped[0]["products_matched"]), 2)
        self.assertIn("超级工厂", deduped[0]["badges"])
        self.assertIn("深度验厂", deduped[0]["badges"])


class TestClassifier(unittest.TestCase):
    def setUp(self):
        self.clf = HeuristicSupplierClassifier()

    def test_factory_classification(self):
        factory_supplier = {
            "supplier_name": "Yongkang Hengtai Hardware & Stainless Steel Manufacturing Co., Ltd. (永康恒泰五金制造有限公司)",
            "business_scope": "Production, custom processing, and R&D for stainless steel vacuum bottles.",
            "badges": ["超级工厂 (Super Factory)", "ISO9001"],
            "factory_area_sqm": 15000,
            "employees_count": 180,
            "registered_capital_k_rmb": 10000,
            "years_in_business": 10
        }
        res = self.clf.classify(factory_supplier)
        self.assertEqual(res.type, "Factory")
        self.assertGreaterEqual(res.confidence, 80.0)
        self.assertTrue(res.is_factory)
        self.assertGreater(len(res.signals), 2)

    def test_trading_company_classification(self):
        trading_supplier = {
            "supplier_name": "Yiwu Haorun International Trading Firm (义乌市浩润商贸商行)",
            "business_scope": "Wholesale distribution and e-commerce retail of general goods.",
            "badges": [],
            "factory_area_sqm": 80,
            "employees_count": 3,
            "registered_capital_k_rmb": 100,
            "years_in_business": 2
        }
        res = self.clf.classify(trading_supplier)
        self.assertEqual(res.type, "Trading Company")
        self.assertGreaterEqual(res.confidence, 70.0)
        self.assertFalse(res.is_factory)

    def test_apify_verified_signals_classification(self):
        supplier = {
            "supplier_name": "Foshan Sanfeng Accessories Co., Ltd.",
            "business_role": "生产加工",
            "is_super_factory": True,
            "is_factory_inspected": True,
            "badges": ["超级工厂 (Super Factory)", "深度验厂 (Factory Inspected)"]
        }
        res = self.clf.classify(supplier)
        self.assertEqual(res.type, "Factory")
        self.assertGreaterEqual(res.confidence, 80.0)


class TestSearchEngine(unittest.TestCase):
    def setUp(self):
        self.engine = SourcingSearchEngine(adapter=DemoAdapter())

    def test_single_product_search(self):
        res = self.engine.search_single_product("Stainless Steel Water Bottle")
        self.assertEqual(res["query"], "Stainless Steel Water Bottle")
        self.assertGreater(res["total_count"], 0)
        self.assertGreaterEqual(res["factories_count"], 1)
        self.assertEqual(res["data_source"], "demo")
        
        first_item = res["results"][0]
        self.assertEqual(first_item["classification"], "Factory")
        self.assertIn("confidence_percent", first_item)
        self.assertIn("shop_url", first_item)

    def test_single_product_factory_only_filter(self):
        res_all = self.engine.search_single_product("Stainless Steel Water Bottle", factory_only=False)
        res_factory_only = self.engine.search_single_product("Stainless Steel Water Bottle", factory_only=True)
        self.assertEqual(res_factory_only["total_count"], res_all["factories_count"])
        for item in res_factory_only["results"]:
            self.assertEqual(item["classification"], "Factory")

    def test_single_search_deduplication(self):
        mock_adapter = MagicMock()
        mock_adapter.get_provider_info.return_value = {"is_demo": False, "name": "Mock Provider"}
        mock_adapter.search_single_product.return_value = [
            {"supplier_id": "b2b-999", "supplier_name": "Dup Factory", "shop_url": "https://shop999.1688.com", "products_matched": ["Item 1"], "primary_product": "Item 1"},
            {"supplier_id": "b2b-999", "supplier_name": "Dup Factory", "shop_url": "https://shop999.1688.com", "products_matched": ["Item 2"], "primary_product": "Item 2"}
        ]
        engine = SourcingSearchEngine(adapter=mock_adapter)
        res = engine.search_single_product("Mug")
        self.assertEqual(res["total_count"], 1)
        self.assertEqual(len(res["results"][0]["matched_products"]), 2)

    def test_multi_product_coverage_search(self):
        queries = [
            "Stainless steel water bottle",
            "Silicone kitchen spatula",
            "Borosilicate glass storage jar",
            "Ceramic coffee mug",
            "Bamboo cutting board"
        ]
        res = self.engine.search_multi_products(queries)
        self.assertEqual(res["total_queries"], 5)
        self.assertGreater(res["total_suppliers_found"], 0)
        self.assertEqual(res["data_source"], "demo")
        
        top_supplier = res["results"][0]
        self.assertGreaterEqual(top_supplier["products_covered_count"], 1)
        self.assertIn("coverage_percent", top_supplier)
        self.assertIn("score", top_supplier)
        self.assertIn("covered_products_list", top_supplier)


class TestExporter(unittest.TestCase):
    def setUp(self):
        self.engine = SourcingSearchEngine(adapter=DemoAdapter())

    def test_single_excel_export(self):
        data = self.engine.search_single_product("Glass Jar")
        excel_io = SourcingDataExporter.export_single_to_excel(data)
        excel_io.seek(0)
        wb = load_workbook(excel_io)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        self.assertEqual(headers[0], "Supplier Name")
        self.assertIn("1688 Product Offer Link", headers)
        self.assertIn("Contact Phone / Mobile", headers)
        self.assertIn("1688 Supplier Contact Page", headers)
        self.assertIn("AliWangWang Web Chat Link", headers)
        self.assertIn("1688 Shop Link", headers)
        self.assertGreater(ws.max_row, 1)

    def test_single_csv_export(self):
        data = self.engine.search_single_product("Glass Jar")
        csv_io = SourcingDataExporter.export_single_to_csv(data)
        content = csv_io.getvalue()
        self.assertTrue(content.startswith(codecs.BOM_UTF8))
        decoded = content.decode("utf-8-sig")
        self.assertIn("Supplier Name", decoded)
        self.assertIn("1688 Product Offer Link", decoded)
        self.assertIn("Contact Phone / Mobile", decoded)
        self.assertIn("1688 Supplier Contact Page", decoded)

    def test_multi_excel_export(self):
        data = self.engine.search_multi_products(["Bottle", "Mug", "Plate"])
        excel_io = SourcingDataExporter.export_multi_to_excel(data)
        excel_io.seek(0)
        wb = load_workbook(excel_io)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        self.assertEqual(headers[0], "Supplier Name")
        self.assertEqual(headers[1], "Products Covered")
        self.assertIn("Matched Product Offer Links", headers)
        self.assertIn("Contact Phone / Mobile", headers)
        self.assertIn("1688 Supplier Contact Page", headers)
        self.assertGreater(ws.max_row, 1)

    def test_multi_csv_export(self):
        data = self.engine.search_multi_products(["Bottle", "Mug", "Plate"])
        csv_io = SourcingDataExporter.export_multi_to_csv(data)
        content = csv_io.getvalue()
        self.assertTrue(content.startswith(codecs.BOM_UTF8))
        decoded = content.decode("utf-8-sig")
        self.assertIn("Products Covered", decoded)
        self.assertIn("Matched Product Offer Links", decoded)
        self.assertIn("Contact Phone / Mobile", decoded)


class TestFastAPIEndpoints(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_get_index(self):
        resp = self.client.get("/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("1688 Sourcing Tool", resp.text)
        self.assertIn("modeToggleBtn", resp.text)

    def test_get_status(self):
        resp = self.client.get("/api/status")
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertIn("demo_mode", data)
        self.assertIn("provider_name", data)

    def test_config_reload_api(self):
        resp = self.client.post("/api/config/reload")
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data["status"], "reloaded")

    def test_toggle_demo_mode_api(self):
        orig_mode = config.demo_mode
        try:
            resp1 = self.client.post("/api/config/toggle-mode", json={"demo_mode": True})
            self.assertEqual(resp1.status_code, 200)
            self.assertTrue(resp1.json()["config"]["demo_mode"])

            resp2 = self.client.post("/api/search/single", json={"product": "Ceramic Mug", "factory_only": False})
            self.assertEqual(resp2.status_code, 200)
            data2 = resp2.json()
            self.assertEqual(data2["query"], "Ceramic Mug")
            self.assertGreater(data2["total_count"], 0)
            self.assertEqual(data2["data_source"], "demo")

            resp3 = self.client.post("/api/search/multi", json={"products": ["Pen", "Notebook", "Eraser"]})
            self.assertEqual(resp3.status_code, 200)
            data3 = resp3.json()
            self.assertEqual(data3["total_queries"], 3)
            self.assertGreater(data3["total_suppliers_found"], 0)
            self.assertEqual(data3["data_source"], "demo")
        finally:
            self.client.post("/api/config/toggle-mode", json={"demo_mode": orig_mode})

    def test_export_single_excel_api(self):
        mock_data = {
            "query": "Tea Kettle",
            "results": [
                {
                    "supplier_name": "Test Factory Co.",
                    "classification": "Factory",
                    "confidence_percent": "95%",
                    "score": 85,
                    "location": "Zhejiang Yongkang",
                    "primary_product": "Tea Kettle",
                    "price_range_rmb": "¥25.00",
                    "moq": "500 pcs",
                    "factory_area_sqm": 12000,
                    "employees_count": 150,
                    "years_in_business": 10,
                    "badges": ["超级工厂"],
                    "signals": ["Verified Factory"],
                    "contact_info": "AliWangWang",
                    "shop_url": "https://shop123.1688.com"
                }
            ]
        }
        export_resp = self.client.post("/api/export/single/excel", json=mock_data)
        self.assertEqual(export_resp.status_code, 200)
        self.assertIn("application/vnd.openxmlformats", export_resp.headers["content-type"])
        self.assertGreater(len(export_resp.content), 1000)

    def test_export_multi_csv_api(self):
        mock_data = {
            "total_queries": 2,
            "results": [
                {
                    "supplier_name": "Test Factory Co.",
                    "products_covered_count": 2,
                    "coverage_percent_str": "100%",
                    "score": 95,
                    "classification": "Factory",
                    "confidence_percent": "95%",
                    "covered_products_list": ["A", "B"],
                    "missing_products_list": [],
                    "location": "Guangdong",
                    "factory_area_sqm": 8000,
                    "employees_count": 100,
                    "years_in_business": 8,
                    "contact_info": "AliWangWang",
                    "shop_url": "https://shop123.1688.com"
                }
            ]
        }
        export_resp = self.client.post("/api/export/multi/csv", json=mock_data)
        self.assertEqual(export_resp.status_code, 200)
        self.assertIn("text/csv", export_resp.headers["content-type"])
        self.assertGreater(len(export_resp.content), 200)

    @patch("requests.post")
    def test_apify_multi_search_mocked_success(self, mock_post):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.json.return_value = [
            {
                "supplierId": "b2b-1001",
                "supplierName": "Yongkang Hengtai Manufacturing",
                "sourceKeywords": ["Bottle", "Mug"],
                "representativeOffers": [{"title": "Insulated Bottle", "priceCny": 15.0, "minimumOrderQuantity": 500}],
                "isSuperFactory": True
            }
        ]
        mock_post.return_value = mock_resp

        adapter = Apify1688Adapter(api_token="apify_api_test")
        results = adapter.search_multi_products(["Bottle", "Mug"])
        self.assertIn("Bottle", results)
        self.assertIn("Mug", results)
        self.assertEqual(len(results["Bottle"]), 1)
        self.assertEqual(results["Bottle"][0]["primary_product"], "Insulated Bottle")

    @patch("requests.post")
    def test_apify_search_http_error_handling(self, mock_post):
        mock_resp = MagicMock()
        mock_resp.status_code = 401
        mock_resp.text = "Unauthorized API Key"
        mock_resp.raise_for_status.side_effect = Exception("401 Client Error: Unauthorized")
        mock_post.return_value = mock_resp

        adapter = Apify1688Adapter(api_token="apify_api_invalid")
        with self.assertRaises(Exception):
            adapter.search_single_product("Tumbler")


    

    def test_single_search_with_custom_pages(self):
        engine = SourcingSearchEngine(adapter=DemoAdapter())
        res = engine.search_single_product("Water Bottle", pages=2)
        self.assertEqual(res["pages_requested"], 2)
        self.assertGreaterEqual(res["total_count"], 10)

    def test_multi_search_with_custom_pages(self):
        engine = SourcingSearchEngine(adapter=DemoAdapter())
        res = engine.search_multi_products(["Water Bottle", "Coffee Mug"], pages=2)
        self.assertEqual(res["pages_requested"], 2)
        self.assertGreaterEqual(res["total_suppliers_found"], 1)

    def test_update_pages_config_api(self):
        resp = self.client.post("/api/config/pages", json={"pages": 3})
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data["status"], "updated")
        self.assertEqual(data["default_pages_per_keyword"], 3)
        # Restore default
        self.client.post("/api/config/pages", json={"pages": 1})

    @patch("requests.post")
    def test_apify_adapter_custom_pages_payload(self, mock_post):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.json.return_value = []
        mock_post.return_value = mock_resp

        adapter = Apify1688Adapter(api_token="apify_api_test")
        adapter.search_single_product("Tumbler", page=3)

        self.assertTrue(mock_post.called)
        _, kwargs = mock_post.call_args
        payload = kwargs.get("json", {})
        self.assertEqual(payload.get("maxPagesPerKeyword"), 3)
        self.assertEqual(payload.get("maxSuppliersPerKeyword"), 60)

if __name__ == "__main__":
    unittest.main()
